"""The dependency-free offline .xlsx writer must produce files real Excel
tooling accepts, and the reader must read files real tooling produced."""
import json
import os
import subprocess
import tempfile
import openpyxl
from helpers import ROOT

JS = os.path.join(ROOT, "tests", "js", "xlsx_roundtrip.js")


def _node(*args):
    out = subprocess.run(["node", JS, *args], capture_output=True, text=True)
    assert out.returncode == 0, out.stderr
    return out.stdout


def test_written_workbook_opens_in_openpyxl_with_correct_values():
    path = os.path.join(tempfile.mkdtemp(), "scores.xlsx")
    _node("write", path)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    assert ws.title == "Scores  JSS1A"          # illegal sheet-name character replaced
    rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    assert rows[1][:6] == ["admission_no", "student_name", "ca1", "ca2", "exam", "total"]
    assert rows[2][0] == "001" and rows[2][2] == 9 and rows[2][3] == 8.5 and rows[2][5] == 77.5
    assert rows[2][1] == 'Obi Chinedu <b>&</b> "Q"'
    assert rows[3][1].startswith("'=HYPERLINK")   # user text is never a live formula
    assert ws["C4"].value is None                 # blank stays blank
    assert rows[4][1] == "Ünïcödé Ẹ̀kọ́"


def test_reader_reads_a_workbook_saved_by_openpyxl():
    path = os.path.join(tempfile.mkdtemp(), "in.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["admission_no", "first_name", "last_name", "class_name"])
    ws.append(["010", "Fatima", "Bello & Sons", "JSS1A"])
    ws.append([11, "Tunde", None, "JSS1A"])
    ws.cell(row=5, column=3, value="gap-row")       # sparse: rows 4 missing entirely
    wb.save(path)                                    # openpyxl deflate-compresses + uses shared strings
    rows = json.loads(_node("read", path))
    assert rows[0] == ["admission_no", "first_name", "last_name", "class_name"]
    assert rows[1] == ["010", "Fatima", "Bello & Sons", "JSS1A"]
    assert rows[2][0] == "11" and rows[2][1] == "Tunde"
    assert rows[-1] == ["", "", "gap-row"]


def test_reader_rejects_garbage_cleanly():
    path = os.path.join(tempfile.mkdtemp(), "bad.xlsx")
    open(path, "wb").write(b"this is not a zip file at all" * 10)
    out = subprocess.run(["node", JS, "read", path], capture_output=True, text=True)
    assert out.returncode == 2 and "valid .xlsx" in out.stderr

"""Real Chromium, real service worker, real IndexedDB: enrol a device, go
offline, and use the offline app end to end (results, printing views, CA3
score entry with validation, CSV/Excel import + export, saved learning
materials), then reconnect and check everything reached the server.
Skipped automatically if Playwright/Chromium isn't installed."""
import io
import json
import os
import re
import tempfile
from urllib.parse import urlparse
import threading

import openpyxl
from werkzeug.serving import make_server
from helpers import fresh_app, ROOT, wait_until

try:
    from playwright.sync_api import sync_playwright
except Exception:          # pragma: no cover
    sync_playwright = None


def _seed(m):
    import db
    conn = db.get_db()
    ids = {}
    ids["school"] = conn.execute("SELECT id FROM schools").fetchone()[0]
    ids["cls"] = conn.execute("SELECT id FROM classes").fetchone()[0]
    ids["term"] = conn.execute("SELECT id FROM terms").fetchone()[0]
    ids["session"] = conn.execute("SELECT id FROM sessions").fetchone()[0]
    subjects = [r[0] for r in conn.execute("SELECT id FROM subjects ORDER BY name")]
    ids["subjects"] = subjects
    students = [r[0] for r in conn.execute("SELECT id FROM students ORDER BY id")]
    conn.execute("UPDATE grading_config SET ca1_max=10, ca2_max=10, ca3_max=10, exam_max=70")
    # a hostile student name: must only ever be shown as text
    conn.execute("UPDATE students SET last_name=? WHERE id=?", ('<img src=x onerror="window.__xss=1">', students[2]))
    data = {students[0]: (9, 8, 7, 60), students[1]: (5, 5, 5, 30), students[2]: (10, 10, 10, 70)}
    for stu, (a, b, c, e) in data.items():
        conn.execute("INSERT INTO scores (student_id, subject_id, term_id, ca1, ca2, ca3, exam) VALUES (?,?,?,?,?,?,?)",
                     (stu, subjects[0], ids["term"], a, b, c, e))
    # a saved-able learning material
    import app as appmod
    school_dir = os.path.join(appmod.MATERIALS_DIR, str(ids["school"]))
    os.makedirs(school_dir, exist_ok=True)
    open(os.path.join(school_dir, "notes.pdf"), "wb").write(b"%PDF-1.4 fake test material")
    conn.execute("INSERT INTO materials (school_id, session_id, class_id, subject_id, title, kind, filename, original_filename) "
                 "VALUES (?,?,?,?,?,?,?,?)",
                 (ids["school"], ids["session"], ids["cls"], subjects[0], "Week 3 Notes", "Notes", "notes.pdf", "notes.pdf"))
    conn.commit()
    ids["students"] = students
    conn.close()
    return ids


def test_offline_app_in_a_real_browser():
    if sync_playwright is None:
        return
    m, _ = fresh_app()
    ids = _seed(m)
    from helpers import LiveServer
    server = LiveServer(m.app).up()
    base = server.base
    problems = []
    typed_student_holder = []
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch()
            ctx = browser.new_context(accept_downloads=True)
            page = ctx.new_page()
            page.on("pageerror", lambda e: problems.append(f"pageerror: {e}"))
            if os.environ.get("DEBUG_E2E"):
                page.on("console", lambda m_: print("CONSOLE", m_.type, m_.text))
                page.on("requestfailed", lambda r_: print("REQFAILED", r_.url))
            dialogs = []
            page.on("dialog", lambda d: (dialogs.append(d.message), d.dismiss()))

            # ---- online: an ordinary login. The login page itself sets up offline
            #      access with the password just typed; there is no separate PIN.
            page.goto(base + "/login")
            page.fill("input[name=username]", "admin")
            page.fill("input[name=password]", "admin123")
            page.click("button[type=submit]")
            page.wait_for_url(re.compile(r".*/app$"))
            page.evaluate("navigator.serviceWorker.ready.then(() => true)")
            # the dashboard fetches the school's data in the background
            wait_until(page, f"OfflineDB.getMeta({ids['school']}, 'last_sync_at').then(v => !!v)", timeout=30.0)
            accts = page.evaluate("OfflineAuth.listAccounts().then(a => a.map(x => ({id: x.device_id, label: x.label, user: x.username, uid: x.user_id})))")
            assert len(accts) == 1, accts
            page.goto(f"{base}/materials?class_id={ids['cls']}&session_id={ids['session']}")
            page.click("button.saveOffline")
            page.wait_for_selector("button.saveOffline:has-text('Saved offline')")
            page.goto(base + "/app")                  # refresh the cached shell like a real user would
            page.wait_for_timeout(300)

            # ---- close the app, lose connectivity, open the app again: it goes
            #      STRAIGHT to the offline app and asks for the password
            page.close()                              # close the app (an unlocked session lives only as long as its tab)
            page = ctx.new_page()
            page.on("pageerror", lambda e: problems.append(f"pageerror: {e}"))
            page.on("dialog", lambda d: (dialogs.append(d.message), d.dismiss()))
            ctx.set_offline(True)
            server.down()                             # really no network: connections are refused
            page.goto(base + "/dashboard")            # what the installed app opens
            assert urlparse(page.url).path == "/app", page.url
            page.wait_for_selector("#pwInput")        # one account on this device -> no account list
            assert "Administrator" in page.locator("h2").first.inner_text()
            page.fill("#pwInput", "admin123")
            page.click("#unlockBtn")
            page.wait_for_selector("h1:has-text('Admin Dashboard')")
            assert page.evaluate("navigator.onLine") is False

            # ---- results & broadsheet, computed on the device
            page.click("#navbarLinks >> text=Classes")
            page.click("#rsLoad")
            page.wait_for_selector("text=Result sheet")
            rows = page.locator("#rsOut table tbody tr")
            assert rows.count() == 3
            first = rows.nth(0).inner_text()
            assert "100" in first                         # 10+10+10+70 tops the class
            assert page.evaluate("window.__xss") is None  # hostile name never executed
            assert "<img" in page.locator("#rsOut").inner_text()   # ...it is shown as plain text
            rows.nth(0).locator("button").click()
            page.wait_for_selector(".sheetView iframe")
            frame = page.frame_locator(".sheetView iframe")
            body_text = frame.locator("body").inner_text()
            assert "CA3" in body_text and "Terminal Report Sheet" in body_text and "100" in body_text
            for f in page.frames:
                assert f.evaluate("window.__xss") is None
            page.click("#rsBroad")
            page.wait_for_selector("text=Broadsheet >> nth=0")
            with page.expect_download() as dl:
                page.click("#rsXlsx")
            path = os.path.join(tempfile.mkdtemp(), "broadsheet.xlsx")
            dl.value.save_as(path)
            ws = openpyxl.load_workbook(path).active
            assert [c.value for c in next(ws.iter_rows())][:2] == ["Pos.", "Student"]

            # ---- score entry: CA3 column, max-mark validation, only changed rows are queued
            page.click("#navbarLinks >> text=Dashboard")
            page.click("#dashWork >> text=Score Entry")
            page.select_option("#seClassSelect", str(ids["cls"]))
            page.wait_for_function("document.querySelectorAll('#seSubjectSelect option').length > 1")
            page.select_option("#seSubjectSelect", str(ids["subjects"][1]))
            page.wait_for_selector("th:has-text('CA3')")
            typed_student = int(page.locator("tr[data-student]").nth(0).get_attribute("data-student"))
            typed_student_holder.append(typed_student)
            inputs = page.locator("tr[data-student] input.ca3")
            inputs.nth(0).fill("99")
            page.click("#saveScoresBtn")
            assert "Nothing was saved" in page.locator("#scoreMsg").inner_text()
            inputs.nth(0).fill("7")
            page.locator("tr[data-student] input.ca1").nth(0).fill("9")
            page.locator("tr[data-student] input.exam").nth(0).fill("50")
            page.click("#saveScoresBtn")
            page.wait_for_function("document.querySelector('#scoreMsg').textContent.includes('Saved 1 score row')", timeout=10000)

            # ---- import a scores workbook (Excel) into another subject
            page.click("#navbarLinks >> text=Dashboard")
            page.click("#navbarLinks >> text=Setup"); page.click("text=Import / Export (CSV & Excel)")
            page.select_option("#ieScClass", str(ids["cls"]))
            page.wait_for_function("document.querySelectorAll('#ieScSubject option').length > 1")
            page.select_option("#ieScSubject", str(ids["subjects"][2]))
            with page.expect_download() as dl:
                page.click("#ieScCsv")
            csv_text = open(dl.value.path(), encoding="utf-8-sig").read()
            assert "admission_no,student_name,ca1,ca2,ca3,exam,total" in csv_text
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["admission_no", "student_name", "ca1", "ca2", "ca3", "exam"])
            ws.append(["001", "x", 8, 8, 8, 60])
            ws.append(["002", "y", 8, 8, 80, 60])       # CA3 out of range -> skipped
            xl = os.path.join(tempfile.mkdtemp(), "scores.xlsx")
            wb.save(xl)
            page.set_input_files("#ieScFile", xl)
            page.wait_for_selector("text=1 student score row(s) ready")
            assert "1 row(s) skipped" in page.locator("#ieScPlan").inner_text()
            page.click("text=Import scores")
            page.wait_for_selector("text=Imported 1 changed row")

            # ---- saved learning material opens offline
            page.click("#navbarLinks >> text=Dashboard")
            page.click("#navbarLinks >> text=Materials")
            page.wait_for_selector("text=Week 3 Notes")
            with page.expect_download() as dl:
                page.click("button:has-text('Open')")
            assert open(dl.value.path(), "rb").read().startswith(b"%PDF-1.4 fake test material")

            # ---- add a staff member offline: only a hash is stored/queued, never the password
            page.click("#navbarLinks >> text=Dashboard")
            page.click("#dashSetup a:has-text(\"Teachers\")")
            page.fill("#tName", "Offline Teacher")
            page.fill("#tUsername", "offlineteacher")
            page.fill("#tPassword", "Sup3r-secret-pw")
            page.click("#tSaveBtn")
            page.wait_for_function("document.querySelector('#tMsg').textContent.includes('Saved') && document.querySelector('#tMsg').textContent.includes('offline')", timeout=30000)
            stored = page.evaluate(f"OfflineDB.getByStatus({ids['school']}, 'users', 'pending').then(r => JSON.stringify(r))")
            assert "Sup3r-secret-pw" not in stored, "plaintext password reached IndexedDB"
            assert "pbkdf2:sha256:600000$" in stored

            # ---- reset an existing teacher's password offline (hashed on the device)
            page.locator("tr:has-text('aokafor') button:has-text('Edit')").click()
            page.fill("#eName", "Mrs Okafor (renamed)")
            page.fill("#ePassword", "Brand-new-pw-1")
            page.click("#eSave")
            page.wait_for_selector("td:has-text('Mrs Okafor (renamed)')")
            stored = page.evaluate(f"OfflineDB.getByStatus({ids['school']}, 'users', 'pending').then(r => JSON.stringify(r))")
            assert "Brand-new-pw-1" not in stored

            # ---- settings: an out-of-range weighting is refused on the device before it ever queues
            page.click("#navbarLinks >> text=Dashboard")
            page.click("#navbarLinks >> text=Settings")
            page.fill("#wExam", "60")
            page.click("#wSave")
            page.wait_for_function("document.querySelector('#wMsg').textContent.includes('already 70')", timeout=10000)
            page.fill("#wExam", "80")
            page.click("#wSave")
            page.wait_for_function("document.querySelector('#wMsg').textContent.includes('exceed 100')", timeout=10000)

            counts = page.evaluate(f"OfflineDB.getPendingCounts({ids['school']})")
            assert counts["pending"] == 4 and counts["conflict"] == 0, counts

            # ---- back online: everything reaches the server
            ctx.set_offline(False)
            server.up()
            page.click("#navbarLinks >> text=Dashboard")
            page.click("#syncPill")
            page.wait_for_function("navigator.onLine === true")
            page.click("button:has-text('Sync Now') >> nth=-1")
            for _ in range(40):
                counts_now = page.evaluate(f"OfflineDB.getPendingCounts({ids['school']})")
                if counts_now["pending"] == 0:
                    break
                page.wait_for_timeout(500)
            browser_pending = page.evaluate(f"OfflineDB.getPendingCounts({ids['school']})")
            assert browser_pending == {"pending": 0, "conflict": 0, "failed": 0}, browser_pending

            # ---- logging out clears cached pages (shared-device safety) but keeps the app shell
            page.goto(base + "/dashboard")
            page.click("a[href='/logout']")
            page.wait_for_url(re.compile(r".*/login"))
            page.wait_for_timeout(600)
            cached = page.evaluate("""(async () => {
                const c = await caches.open('school-results-shell-v7');
                return (await c.keys()).map(r => new URL(r.url).pathname);
            })()""")
            assert "/dashboard" not in cached and "/materials" not in cached, cached
            assert "/app" in cached and "/static/js/offline-results.js" in cached, cached
            browser.close()
    finally:
        server.down()
    assert not problems, problems

    import db
    conn = db.get_db()
    rows = conn.execute("SELECT student_id, subject_id, ca1, ca2, ca3, exam FROM scores WHERE subject_id IN (?,?) ORDER BY subject_id, student_id",
                        (ids["subjects"][1], ids["subjects"][2])).fetchall()
    got = {(r["subject_id"], r["student_id"]): (r["ca1"], r["ca2"], r["ca3"], r["exam"]) for r in rows}
    if os.environ.get("DEBUG_E2E"):
        print("SCORES", [tuple(r) for r in conn.execute("SELECT * FROM scores")])
        print("AUDIT", [(r["entity"], r["outcome"], r["note"], r["client_ts"]) for r in conn.execute("SELECT * FROM change_audit")])
    stu = ids["students"]
    assert got[(ids["subjects"][1], typed_student_holder[0])] == (9, 0, 7, 50)     # typed in the score screen
    assert got[(ids["subjects"][2], stu[0])] == (8, 8, 8, 60)     # imported from the Excel file
    assert (ids["subjects"][2], stu[1]) not in got                 # the out-of-range row was skipped
    audit = conn.execute("SELECT COUNT(*) FROM change_audit WHERE outcome='applied' AND entity='scores'").fetchone()[0]
    assert audit == 2
    conn.close()
    # the staff account created offline, from a password that never left the browser in plaintext, logs in
    from helpers import login
    assert login(m.app.test_client(), "offlineteacher", "Sup3r-secret-pw").status_code == 302
    assert login(m.app.test_client(), "aokafor", "Brand-new-pw-1").status_code == 302        # reset offline
    assert login(m.app.test_client(), "aokafor", "teacher123").status_code != 302
